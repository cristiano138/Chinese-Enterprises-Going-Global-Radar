from pathlib import Path
from datetime import date
import subprocess, sys
try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', '--quiet', 'openpyxl>=3.1.0'])
    import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

ROOT = Path('E:/workbuddy/2026-10-14-01-02-52/compliance-radar')
OUT = ROOT / 'docs/项目功能清单.xlsx'
assert not OUT.exists(), '禁止覆盖已有工作簿'
def xl_color(value):
    value = value.removeprefix('#').upper()
    assert len(value) == 6
    return 'FF' + value
XL_MAIN = xl_color('#548235')
XL_LIGHT = xl_color('#E2EFDA')
XL_WHITE = xl_color('#FFFFFF')
XL_TEXT = xl_color('#1A1A1A')
XL_WARN = xl_color('#FFEB9C')
XL_RED = xl_color('#FFC7CE')
wb = openpyxl.Workbook()
wb.remove(wb.active)
wb.properties.title = 'ComplianceRadar 项目功能清单'
wb.properties.description = '2026-10-17 项目审计快照；历史验证与本轮验证明确区分。不是法规法律意见。'

def sheet(name, headers, rows, widths, status_col=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    end = len(rows) + 1
    for cell in ws[1]:
        cell.font = Font(name='微软雅黑', bold=True, color=XL_WHITE)
        cell.fill = PatternFill('solid', fgColor=XL_MAIN)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2, max_row=end):
        for c in row:
            c.font = Font(name='微软雅黑', size=11, color=XL_TEXT)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if c.row % 2 == 0:
                c.fill = PatternFill('solid', fgColor=XL_LIGHT)
            if isinstance(c.value, date):
                c.number_format = 'yyyy-mm-dd'
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'C2' if len(headers) > 6 else 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{end}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    if status_col:
        col = get_column_letter(status_col)
        dv = DataValidation(type='list', formula1='"已实现,部分实现,占位,待开发"')
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{end}')
        for value, color in [('已实现', XL_LIGHT), ('部分实现', XL_WARN), ('占位', XL_RED), ('待开发', XL_RED)]:
            ws.conditional_formatting.add(f'{col}2:{col}{end}', CellIsRule(operator='equal', formula=['"' + value + '"'], fill=PatternFill('solid', fgColor=color)))
    return ws

sheet('阅读说明', ['项目', '内容', '依据或注意事项'], [
    ['项目定位', '出海合规信息查询、变化整理与 OPC 上线准备工作台', '不是法律意见、认证结论或自动合规裁定'],
    ['版本日期', date(2026,10,17), '本表为当前审计快照，不代表此后自动联网更新'],
    ['配套文件', '项目总说明书.md；项目更新记录.md', '均位于 docs/；以后每次更新同步维护三份文件'],
    ['当前市场', '欧盟、美国', '五个品类入口；不代表覆盖全部法规'],
    ['数据边界', '86 条候选；29 条已核对；57 条待核对', 'release_eligible=false；不可宣称全部法规已核对'],
    ['验证口径', '历史测试证据不等于本轮重跑', '本轮仅核验文档、代码片段、报告和工作簿结构'],
    ['上线口径', '历史 ModelScope 新版已验收；GitHub 最新同步待核验', '本轮文档新增尚未推送或重新部署'],
    ['模块命名', 'M5 区域对比；M6 多区域汇总', '按代码模块命名；历史页面导航编号存在反向命名记录，应统一'],
], [22,65,75])

features = [
['F01','M1','合规地图','已实现','单品类、单市场','法规列表、编号、链接、核对状态','modules/m1_map.py','历史 M1 专项与统一门禁通过','仅本地候选语料，不保证法规穷尽','保持数据核验'],
['F02','M1.5','自然语言路由','已实现','业务描述与上下文','单一品类和市场，复用 M1','modules/m15_router.py','历史路由、降级、缓存与转义测试通过','各取前500字符，不支持矩阵或文件解析','扩展多槽位前先做契约测试'],
['F03','M2','法规变化时间线','已实现','品类、市场、时间范围','法规事件、证据及参考事件','modules/m2_timeline.py; core/dates.py','历史 M2/M3 专项与统一门禁通过','首次收录不是生效；美国日期仍有缺口','逐条核实法律日期'],
['F04','M3','合规待办','已实现','已人工核对法规','动作、条件、期限、稳定 action_id','modules/m3_todo.py; core/actions.py','历史待办与门禁测试通过','未核对美国法规不能进入执行层','完善真实场景验证'],
['F05','M3','待办 CSV 导出','已实现','有效动作列表或动作ID','带官方证据的 CSV','modules/m3_todo.py','历史空导出、未知ID测试通过','不是 PDF 或正式法律文件','保持导出契约'],
['F06','M4','模型对照回放','占位','固定玩具×欧盟题干','回放入口或未生成提示','modules/m4_compare.py','历史审计确认无完整真实回放','不能把示意当作实验结果','保存真实模型输出、时间与配置'],
['F07','M5','同品类区域对比','已实现','单品类、多市场','共同编号与区域独有法规','modules/m5_compare.py','历史 M5/M6 专项通过','编号完全匹配不等于法律等效','完善义务维度比较'],
['F08','M6','多区域召回汇总','已实现','单品类、多个市场','按区域统计法规与核对状态','modules/m6_multi_region.py','历史空选择与部分失败测试通过','支持多条结构化记录，不是多份用户文件解析','保留区域失败传播'],
['F09','M7','OPC 上线准备','已实现','M3 已核对动作','六阶段、角色、优先级、证据','modules/m7_opc_readiness.py','历史 M7 专项与统一门禁通过','不判定商品已合规；美国区域可能部分可用','商家试用验证'],
['F10','M7','OPC CSV 导出','已实现','上线准备任务','UTF-8 BOM 执行清单','modules/m7_opc_readiness.py; app_main.py','历史 M7 导出测试通过','不含多人协作或状态持久化','保持导出与页面一致'],
['F11','安全','链接与引用校验','已实现','法规记录、模型引用','白名单链接与拦截提示','core/verify.py','历史 URL 绕过与引用契约测试通过','域名合法不证明内容最新或正文可访问','补充正文核验'],
['F12','数据','人工核对门禁','已实现','法规人工核对状态','M3/M7 执行层过滤','core/verify.py; modules/m3_todo.py','历史发布门禁测试通过','结构 valid=true 不等于 release_eligible=true','核对美国57条'],
['F13','模型','DeepSeek 兼容调用','已实现','三项环境变量和提示','模型理解、润色与规则降级','core/llm.py','历史配置与降级路径检查','本轮未改 DeepSeek 配置；密钥不入库','继续平台 Secret 注入'],
['F14','性能','路由缓存及超时','已实现','业务输入哈希','128项内存缓存、12秒超时兜底','modules/m15_router.py','历史缓存超时专项通过','超时不代表底层 HTTP 已取消','并发资源压力验证'],
['F15','性能','通用磁盘缓存治理','部分实现','模型请求','data/cache 磁盘缓存','core/llm.py','历史代码审计','未实现明确 TTL 和容量淘汰','增加过期和限容策略'],
['F16','界面','新版工作台与动画','已实现','页面访问、模块切换','淡黄色、左导航、响应式和开场动画','app_main.py','历史线上 DOM 与截图验收','未完成完整多浏览器回归；页面与代码模块编号待统一','浏览器交互回归'],
['F17','界面','首页数据准确标注','部分实现','快照统计','主文案候选语料；卡片仍标已核对法规','app_main.py:_overview_html','本轮检查确认概览卡片仍为旧口径','86条不能称为已核对法规','拆分候选、已核对和待核对数'],
['F18','部署','ModelScope 创空间','已实现','完整运行文件及平台 Secret','Gradio 6.2.0 运行实例','app.py; ms_deploy.json; requirements.txt','历史 d4cb4f2 部署与浏览器验收','本轮新增文档未发布；Running不等于新版验收','更新后核验提交与真实页面'],
['F19','发布','GitHub 最新同步','部分实现','本地 Git 提交','远端仓库版本','Git origin','历史 TLS/代理失败；最新状态未核验','不能将本地提交当作远端成功','恢复连接后核验远端提交'],
['F20','扩展','多品类×多市场矩阵','待开发','多个品类和市场','组合结果及覆盖状态','当前无完整实现','历史接口审计','当前路由仅单值槽位','新增矩阵契约与限流'],
['F21','扩展','文件上传与长文解析','待开发','PDF、Word、BOM或合同','解析、分块、证据和检索','当前未实现','历史代码审计','当前不能上传或完整分析用户文档','先证据对象，再解析和异步任务'],
['F22','扩展','账号权限和团队协作','待开发','成员、角色、任务状态','持久化协作工作台','当前未实现','历史架构审计','无账号、数据库、权限与长期任务状态','按实际商业需求实施'],
['F23','扩展','全文快照与订阅推送','待开发','法规来源与订阅条件','版本差异、更新通知','当前无完整实现','历史能力审计','链接检查不是法规全文实时监控','补齐快照和变化验证'],
]
sheet('功能清单', ['功能ID','模块','功能名称','状态','输入','输出或实现','代码位置','验证证据','限制与风险','后续动作'], features, [12,14,24,15,32,44,48,44,50,42],4)

sheet('数据状态', ['指标','记录值','来源','解释与风险'], [
['候选语料总数',86,'data/data_quality_report.json','候选不等于已人工核对'],
['唯一键',86,'data/data_quality_report.json','重复键0'],
['欧盟记录数',29,'data/snapshot.json','已人工核对；其中含已废止记录，不能都当现行法规'],
['美国记录数',57,'data/snapshot.json','待人工核对'],
['已人工核对',29,'audit_status_distribution','不得改写成86条已核对'],
['待人工核对',57,'audit_status_distribution','不得进入 M3/M7 执行层'],
['缺失官方链接',0,'data/data_quality_report.json','有链接不代表正文已核验'],
['结构校验','valid=true','data/data_quality_report.json','结构有效，不代表法律准确'],
['发布资格','release_eligible=false','data/data_quality_report.json','完整数据发布门禁未满足'],
['现行有效',25,'status_distribution','报告快照，不是本轮法律复核'],
['已废止',4,'status_distribution','应保留效力状态与替代依据'],
['持续修订',57,'status_distribution','美国数据；需核实日期和现行文本'],
['数据截止','2026.10.14','data/snapshot.json','不代表每天自动更新'],
['链接时效报告','28 ok / 58 net_error','历史 freshness 审计','202挑战页与网络错误不能直接判有效或失效；115/86链接口径待统一'],
['美国补充方向','MoCRA、Prop 65、CPSA及真实生效日期','docs/问题与备忘/待补充法规清单.md','缺口记录不是法律完整性保证'],
], [24,34,50,75])

sheet('验证记录', ['记录ID','范围','结果','证据','验证阶段','未验证事项'], [
['V01','M1/M1.5/M2/M3/M5/M6/M7及输出契约','历史通过，失败数0','scripts/selftest_all.py','前序开发阶段','本轮未重跑应用测试'],
['V02','M1.5缓存、超时、降级、转义','历史通过','scripts/selftest_m15.py','前序开发阶段','真实高并发未测'],
['V03','M5/M6空选择及区域失败','历史通过','scripts/selftest_m5_m6.py','前序开发阶段','多品类矩阵未实现'],
['V04','M7门禁及导出','历史通过','scripts/selftest_m7.py','前序开发阶段','真实商家长期试用未完成'],
['V05','Gradio 6.2.0与新版线上页面','历史验收通过','ModelScope d4cb4f2；线上新版页面.png','前序部署阶段','本轮无重新部署'],
['V06','数据报告及首页卡片口径','本轮已核验','data_quality_report.json；app_main.py:1096','本轮文档整理','法律内容未重新人工审查；首页卡片待修复'],
['V07','Excel结构和关键记录','生成后脚本校验','docs/.项目功能清单.ref/build.py','本轮文档整理','不是应用运行或法律数据验收'],
], [12,42,30,62,24,70])

sheet('更新维护规则', ['序号','触发或对象','必须执行','记录位置','验收条件'], [
[1,'代码或功能变化','同步实现状态、输入输出、代码位置、限制','功能清单；项目总说明书','不将路线图写成已实现'],
[2,'数据变化','更新候选、已核对、待核对及发布资格','数据状态；说明书数据章节','数量与最新报告一致，注明快照日期'],
[3,'测试变化','记录实际运行命令、时间、结果及未测项','验证记录；项目更新记录','历史通过不冒充本次重跑'],
[4,'部署变化','分别记录 GitHub 与 ModelScope 提交和运行验收','项目更新记录','仓库提交、构建和页面三层核验；密钥不入库'],
[5,'每次完成更新','追加日期、版本、文件、改动、验证、风险、同步状态','docs/项目更新记录.md','追加历史，不覆盖旧记录'],
[6,'三件套联动','同轮维护总说明书、功能清单和更新记录','docs/','未能验证的项目明确标注待核验'],
[7,'后续会话恢复','先读三件套与项目记忆再修改','项目文档及 .workbuddy/memory/','这是维护约定，不是后台自动同步服务'],
[8,'本轮交付','新建Excel并核验关键行；补充更新日志','项目功能清单.xlsx','本轮不修改应用代码、法规底表或模型配置'],
], [10,26,65,50,65])
wb.save(OUT)
check = openpyxl.load_workbook(OUT)
assert check.sheetnames == ['阅读说明','功能清单','数据状态','验证记录','更新维护规则']
assert check['功能清单'].max_row == len(features) + 1
assert check['功能清单']['D7'].value == '占位'
assert check['数据状态']['B2'].value == 86
assert check['数据状态']['B6'].value == 29
assert check['数据状态']['B7'].value == 57
assert not any(c.data_type == 'f' for ws in check for row in ws for c in row), '本快照表不应包含公式'
print('PASS: 5 sheets; 23 feature records; counts 86/29/57; no formulas; ' + str(OUT))
